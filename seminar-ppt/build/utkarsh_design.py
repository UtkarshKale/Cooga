"""Utkarsh's design system: "neon night".

Dark night-blue canvas with soft violet/cyan glows, translucent glass panels, violet→cyan
gradient accents, Calibri type with Consolas code-style section labels. Built to look nothing
like Vedika's white amber/charcoal "hazard signage" deck. Only the shape/text primitives are
shared with design.py.
"""
import io
import os

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from design import arrow, box, fill, line, notes, path, rgb, text  # noqa: F401

BODY = "Calibri"
MONO = "Consolas"
BG = "0B1020"
TEXT = "EEF2FF"
MUTED = "A3ADC8"
DIM = "6B7699"
PANEL = "151C38"
PANEL2 = "1B2445"
EDGE = "2C3870"
CYAN = "22D3EE"
VIOLET = "A78BFA"
DEEP = "7C3AED"
PINK = "F472B6"
GREEN = "34D399"
RED = "F87171"
CODE_BG = "070B17"
WHITE = "FFFFFF"
ICON = {CYAN: "uc", VIOLET: "uv", PINK: "up", GREEN: "ug", RED: "ur", MUTED: "us", BG: "ud", WHITE: "w"}
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def t(s, x, y, w, h, paras, size=14, color=TEXT, **kw):
    kw.setdefault("font", BODY)
    return text(s, x, y, w, h, paras, size=size, color=color, **kw)


def label(s, x, y, w, txt, color=CYAN, size=10.5, align="l"):
    return t(s, x, y, w, 0.28, txt, size=size, color=color, font=MONO, bold=True, align=align)


def glass(s, x, y, w, h, bg=PANEL, edge=EDGE, radius=0.14, alpha=90):
    return box(s, x, y, w, h, bg=bg, line=edge, lw=0.75, radius=radius, alpha=alpha)


def gradient(shape, c1=DEEP, c2="0891B2", angle=0):
    f = shape.fill
    f.gradient()
    f.gradient_angle = angle
    st = f.gradient_stops
    st[0].position, st[1].position = 0.0, 1.0
    st[0].color.rgb = rgb(c1)
    st[1].color.rgb = rgb(c2)
    return shape


def grad_box(s, x, y, w, h, paras=None, size=14, radius=0.14, shape="rrect", c1=DEEP, c2="0891B2", angle=0, **kw):
    shp = box(s, x, y, w, h, bg=c1, radius=radius, shape=shape, paras=paras, size=size, color=WHITE, **kw)
    return gradient(shp, c1, c2, angle)


class UDeck:
    def __init__(self, template, icon_dir, bg_dir, footer, total):
        self.icon_dir, self.bg_dir, self.footer, self.total = icon_dir, bg_dir, footer, total
        self.prs = Presentation(template)
        self.logo = self.banner = None
        for sl in self.prs.slides:
            for shp in sl.shapes:
                if shp.shape_type == 13:
                    if shp.image.size == (1100, 140):
                        self.banner = shp.image.blob
                    elif shp.image.size == (316, 238):
                        self.logo = shp.image
        lst = self.prs.slides._sldIdLst
        for sldId in list(lst):
            self.prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
        self.blank = self.prs.slide_masters[0].slide_layouts[6]

    def save(self, p):
        self.prs.save(p)

    # ------------------------------------------------ primitives
    def icon(self, s, name, x, y, d, color=CYAN):
        return s.shapes.add_picture(os.path.join(self.icon_dir, f"{name}_{ICON[color]}.png"),
                                    Inches(x), Inches(y), Inches(d), Inches(d))

    def icon_orb(self, s, name, x, y, d, c1=DEEP, c2="0891B2"):
        grad_box(s, x, y, d, d, shape="oval", c1=c1, c2=c2, angle=45)
        pad = d * 0.25
        self.icon(s, name, x + pad, y + pad, d - 2 * pad, WHITE)

    def logo_pic(self, s, x, y, w):
        h = w * self.logo.size[1] / self.logo.size[0]
        s.shapes.add_picture(io.BytesIO(self.logo.blob), Inches(x), Inches(y), Inches(w), Inches(h))

    def _slide(self, bg="bg_content.jpg"):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BG)
        s.shapes.add_picture(os.path.join(self.bg_dir, bg), 0, 0, Inches(10), Inches(7.5))
        return s

    def _footer(self, s):
        t(s, 0.6, 7.02, 6.5, 0.28, self.footer, size=9, color=DIM, font=MONO)
        tb = s.shapes.add_textbox(Inches(8.2), Inches(7.02), Inches(1.2), Inches(0.28))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn("a:fld"), id="{5E2A9C41-7B3D-4F6E-8A1B-C2D3E4F5A6B7}", type="slidenum")
        rPr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz="900", b="1")
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=CYAN)
        etree.SubElement(rPr, qn("a:latin"), typeface=MONO)
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))
        r = p.add_run()
        r.text = f" / {self.total}"
        r.font.size, r.font.name = Pt(9), MONO
        r.font.color.rgb = rgb(DIM)

    def _logo_badge(self, s, x=8.82, y=0.32, d=0.66):
        box(s, x, y, d, d, bg=WHITE, radius=0.14)
        self.logo_pic(s, x + 0.05, y + 0.08, d - 0.1)

    # ------------------------------------------------ frames
    def slide(self, tag, title, sub=None, size=None):
        s = self._slide()
        label(s, 0.6, 0.4, 8.0, tag)
        if size is None:
            size = 28 if len(title) <= 38 else 24
        t(s, 0.6, 0.66, 8.0, 0.62, title, size=size, bold=True, anchor="m")
        if sub:
            t(s, 0.6, 1.25, 8.6, 0.32, sub, size=12.5, color=MUTED, italic=True)
        self._logo_badge(s)
        self._footer(s)
        return s

    def institution_card(self, s):
        box(s, 0.4, 0.3, 9.2, 1.35, bg=WHITE, radius=0.16)
        self.logo_pic(s, 0.6, 0.38, 1.25)
        pic = s.shapes.add_picture(io.BytesIO(self.banner), Inches(1.95), Inches(0.4), Inches(7.45), Inches(1.11))
        pic.crop_left, pic.crop_right = 0.11745, 0.03185

    def orbit(self, s, cx, cy, r=1.35, core="LLM"):
        """Agent orbit motif: LLM core with planning / tools / memory / reflection satellites."""
        box(s, cx - r, cy - r, 2 * r, 2 * r, shape="oval", line=VIOLET, lw=1.25, line_alpha=45)
        box(s, cx - r * 0.62, cy - r * 0.62, 2 * r * 0.62, 2 * r * 0.62, shape="oval", line=CYAN, lw=1.0,
            line_alpha=35)
        box(s, cx - 0.62, cy - 0.62, 1.24, 1.24, bg=DEEP, shape="oval", alpha=30)
        grad_box(s, cx - 0.45, cy - 0.45, 0.9, 0.9, shape="oval", paras=[[(core, {"bold": True})]], size=15,
                 angle=45, font=MONO)
        sats = [(-90, "plan", "Plan"), (0, "tools", "Tools"), (90, "db", "Memory"), (180, "sync", "Reflect")]
        import math
        for ang, ic, lab in sats:
            a = math.radians(ang)
            x, y = cx + r * math.cos(a), cy + r * math.sin(a)
            box(s, x - 0.3, y - 0.3, 0.6, 0.6, bg=PANEL2, line=CYAN, lw=1.25, shape="oval")
            self.icon(s, ic, x - 0.16, y - 0.16, 0.32, CYAN)
            if ang == -90:
                t(s, x + 0.36, y - 0.13, 1.0, 0.26, lab, size=10.5, color=MUTED, font=MONO)
            else:
                t(s, x - 0.6, y + 0.33, 1.2, 0.26, lab, size=10.5, color=MUTED, align="c", font=MONO)

    def code(self, s, x, y, w, h, lines, size=11.5, title="pseudocode"):
        glass(s, x, y, w, h, bg=CODE_BG, edge=EDGE, radius=0.12)
        for i, c in enumerate(("FF5F57", "FEBC2E", "28C840")):
            box(s, x + 0.2 + i * 0.2, y + 0.2, 0.12, 0.12, bg=c, shape="oval")
        t(s, x + 0.85, y + 0.12, w - 1.0, 0.28, title, size=10, color=DIM, font=MONO)
        paras = []
        for i, ln in enumerate(lines, 1):
            code_part, _, comment = ln.partition("#")
            runs = [(f"{i:>2}  ", {"color": "475079"}), (code_part, {"color": "E2E8F0"})]
            if comment:
                runs.append(("# " + comment.strip(), {"color": VIOLET, "italic": True}))
            paras.append(runs)
        text(s, x + 0.2, y + 0.5, w - 0.3, h - 0.55, paras, size=size, font=MONO, after=2)

    def table(self, s, rows, widths, y=1.65, header_h=0.45, row_h=0.6, size=12, num_col=False, hl_col=None):
        nr, nc = len(rows), len(widths)
        gf = s.shapes.add_table(nr, nc, Inches(0.6), Inches(y), Inches(sum(widths)),
                                Inches(header_h + row_h * (nr - 1)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for k in ("bandRow", "firstRow"):
            tblPr.set(k, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
        sid.text = NO_STYLE_TABLE
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(nr):
            tbl.rows[i].height = Inches(header_h if i == 0 else row_h)
            for j in range(nc):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.1)
                c.margin_top = c.margin_bottom = Inches(0.05)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid()
                if i == 0:
                    bg = DEEP if j != hl_col else "0891B2"
                elif j == hl_col:
                    bg = "123047"
                else:
                    bg = PANEL if i % 2 else PANEL2
                c.fill.fore_color.rgb = rgb(bg)
                tf = c.text_frame
                tf.word_wrap = True
                v = rows[i][j]
                if i == 0:
                    fill(tf, v, 12, WHITE, True, font=BODY)
                elif num_col and j == 0:
                    fill(tf, v, 14, CYAN, True, font=MONO, align="c")
                else:
                    fill(tf, v, size, TEXT, j == 0 and not num_col, font=BODY, after=2)
        return tbl
