"""Design system for the redesigned seminar decks.

Editorial layout on the college template's 4:3 canvas: dark title / conclusion / closing
slides around light content slides, a section "eyebrow" above a message-style title,
one dominant colour per deck plus one sharp accent, and native (editable) diagrams.
The institute logo and header banner are taken from the BIT Seminar-I template.
"""
import io
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HEAD = "Cambria"
BODY = "Calibri"
MONO = "Consolas"

GOOD = "2E9E6A"
BAD = "D64545"
AMBER = "F2C94C"
WHITE = "FFFFFF"

AGENTIC = dict(
    ink="17153A", primary="5B4CDB", accent="F2643D", soft="EEEDFC", card="F5F5FA",
    muted="6A6886", line="DCDAF0", on_dark="B4B0E6", dark_line="34315E", text="1F1D33",
    code_num="8A86B8", ic_p="ap", ic_a="aa", ic_i="ai",
)
DISASTER = dict(
    ink="0E2B31", primary="12807F", accent="EA6A1E", soft="E3F1F0", card="F3F7F7",
    muted="5B6E72", line="CBE0DE", on_dark="9CCBC7", dark_line="2A4B51", text="17282C",
    code_num="7FA9A6", ic_p="dp", ic_a="da", ic_i="di",
)

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rrect": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
          "diamond": MSO_SHAPE.DIAMOND, "chevron": MSO_SHAPE.CHEVRON}


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- text
def _run(p, text, size, color, bold=False, italic=False, font=BODY, spc=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.color.rgb = rgb(color)
    f.bold = bold
    f.italic = italic
    f.name = font
    if spc:
        r.font._rPr.set("spc", str(spc))
    return r


def _bullet(p, char, color, indent):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    bc = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bc, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def fill(tf, paras, size=14, color="000000", bold=False, italic=False, align="l", font=BODY,
         after=4, bullet=None, bullet_color=None, indent=0.2, spc=None, line=None):
    """paras: str | list of items; item = str | [(text, opts), ...] | {"runs":..., overrides}."""
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        o = {}
        if isinstance(item, dict):
            o, runs = item, item["runs"]
        else:
            runs = item
        if isinstance(runs, str):
            runs = [(runs, {})]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[o.get("align", align)]
        p.space_after = Pt(o.get("after", after))
        if line:
            p.line_spacing = line
        if o.get("bullet", bullet):
            _bullet(p, o.get("bullet", bullet), o.get("bullet_color", bullet_color or color), indent)
        for t, ro in runs:
            _run(p, t, ro.get("size", o.get("size", size)), ro.get("color", o.get("color", color)),
                 ro.get("bold", o.get("bold", bold)), ro.get("italic", o.get("italic", italic)),
                 ro.get("font", o.get("font", font)), ro.get("spc", spc))


def text(s, x, y, w, h, paras, size=14, color="000000", bold=False, italic=False, align="l",
         anchor="t", font=BODY, after=4, margin=0.0, bullet=None, bullet_color=None, indent=0.2,
         spc=None, line=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = ANCHOR[anchor]
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    fill(tf, paras, size, color, bold, italic, align, font, after, bullet, bullet_color, indent, spc, line)
    return tb


# ---------------------------------------------------------------- shapes
def _strip(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _alpha(fill_parent, alpha):
    clr = fill_parent.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 1000)))


def box(s, x, y, w, h, bg=None, line=None, lw=1.0, shape="rrect", radius=0.1, paras=None, size=13,
        color="000000", bold=False, italic=False, align="c", anchor="m", font=BODY, margin=0.06,
        after=0, dash=False, alpha=None, line_alpha=None, spc=None):
    shp = s.shapes.add_shape(SHAPES[shape], Inches(x), Inches(y), Inches(w), Inches(h))
    _strip(shp)
    if shape == "rrect":
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if bg:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(bg)
        if alpha is not None:
            _alpha(shp._element.spPr.find(qn("a:solidFill")), alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = 7
        if line_alpha is not None:
            _alpha(shp._element.spPr.find(qn("a:ln")), line_alpha)
    else:
        shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = ANCHOR[anchor]
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if paras is not None:
        fill(tf, paras, size, color, bold, italic, align, font, after, spc=spc)
    return shp


def _conn(s, x1, y1, x2, y2, color, w, dash=False, alpha=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _strip(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 7
    if alpha is not None:
        _alpha(c.line._get_or_add_ln(), alpha)
    return c


def line(s, x1, y1, x2, y2, color, w=1.0, dash=False, alpha=None):
    return _conn(s, x1, y1, x2, y2, color, w, dash, alpha)


def arrow(s, x1, y1, x2, y2, color, w=1.5, head="end", dash=False):
    c = _conn(s, x1, y1, x2, y2, color, w, dash)
    ln = c.line._get_or_add_ln()
    if head in ("start", "both"):
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if head in ("end", "both"):
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def path(s, pts, color, w=1.5, dash=False, head="end"):
    n = len(pts) - 1
    for i in range(n):
        h = "end" if (i == n - 1 and head == "end") else None
        arrow(s, *pts[i], *pts[i + 1], color, w, head=h, dash=dash)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, template, icon_dir, P, footer):
        self.P = P
        self.footer_text = footer
        self.icon_dir = icon_dir
        self.prs = Presentation(template)
        slides = list(self.prs.slides)
        self.logo = self.banner = None
        for sl in slides:
            for shp in sl.shapes:
                if shp.shape_type == 13:
                    if shp.image.size == (1100, 140):
                        self.banner = shp.image.blob
                    elif shp.image.size == (316, 238):
                        self.logo = shp.image
        lst = self.prs.slides._sldIdLst
        for sldId in list(lst):
            self.prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
        self.blank = self.prs.slide_masters[0].slide_layouts[6]

    def save(self, path):
        self.prs.save(path)

    # -------------------------------------------- primitives bound to the palette
    def icon(self, s, name, x, y, d, key):
        return s.shapes.add_picture(os.path.join(self.icon_dir, f"{name}_{key}.png"),
                                    Inches(x), Inches(y), Inches(d), Inches(d))

    def icon_circle(self, s, name, x, y, d, bg=None, key="w"):
        box(s, x, y, d, d, bg=bg or self.P["primary"], shape="oval")
        pad = d * 0.25
        return self.icon(s, name, x + pad, y + pad, d - 2 * pad, key)

    def logo_pic(self, s, x, y, w):
        h = w * self.logo.size[1] / self.logo.size[0]
        return s.shapes.add_picture(io.BytesIO(self.logo.blob), Inches(x), Inches(y), Inches(w), Inches(h))

    def _page_number(self, s, color):
        tb = s.shapes.add_textbox(Inches(8.4), Inches(7.0), Inches(1.0), Inches(0.3))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn("a:fld"), id="{8A1F7C52-5D3B-4C1E-9E0B-2B7F1C3D4E5F}", type="slidenum")
        rPr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz="900", b="1")
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(rPr, qn("a:latin"), typeface=BODY)
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))

    def _bg(self, s, color):
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(color)

    # -------------------------------------------- frames
    def content(self, eyebrow, title, dark=False, size=None):
        P = self.P
        s = self.prs.slides.add_slide(self.blank)
        self._bg(s, P["ink"] if dark else WHITE)
        text(s, 0.6, 0.42, 7.6, 0.3, eyebrow.upper(), size=10.5, bold=True, color=P["accent"], spc=160)
        if size is None:
            size = 26 if len(title) <= 40 else 23
        text(s, 0.6, 0.7, 7.9, 0.75, title, size=size, bold=True, font=HEAD,
             color=WHITE if dark else P["ink"], anchor="t")
        if dark:
            box(s, 8.72, 0.28, 0.72, 0.72, bg=WHITE, radius=0.12)
            self.logo_pic(s, 8.76, 0.35, 0.64)
        else:
            self.logo_pic(s, 8.7, 0.3, 0.72)
        text(s, 0.6, 7.0, 6.5, 0.3, self.footer_text, size=9, color=P["on_dark"] if dark else P["muted"])
        self._page_number(s, P["on_dark"] if dark else P["muted"])
        return s

    def cover(self, eyebrow, title, subtitle, motif, info=None, title_size=40):
        """Dark cover: white institutional header (logo + BIT banner), dark hero below."""
        P = self.P
        s = self.prs.slides.add_slide(self.blank)
        self._bg(s, P["ink"])
        box(s, 0, 0, 10, 1.5, bg=WHITE, shape="rect")
        self.logo_pic(s, 0.35, 0.2, 1.4)
        pic = s.shapes.add_picture(io.BytesIO(self.banner), Inches(1.85), Inches(0.17), Inches(7.85), Inches(1.17))
        pic.crop_left, pic.crop_right = 0.11745, 0.03185
        motif(self, s)
        text(s, 0.6, 2.0, 6.5, 0.3, eyebrow.upper(), size=11, bold=True, color=P["accent"], spc=200)
        text(s, 0.6, 2.35, 6.3, 1.8, title, size=title_size, bold=True, font=HEAD, color=WHITE, after=0)
        if subtitle:
            text(s, 0.6, 3.45 if title_size >= 40 else 3.3, 5.9, 1.3, subtitle, size=17, color=P["on_dark"], font=HEAD,
                 italic=True)
        if info:
            for k, (label, l1, l2) in enumerate(info):
                x = 0.6 + k * 4.4
                text(s, x, 5.45, 4.2, 0.28, label.upper(), size=10, bold=True, color=P["accent"], spc=180)
                text(s, x, 5.75, 4.2, 0.45, l1, size=19, bold=True, color=WHITE, font=HEAD)
                text(s, x, 6.2, 4.2, 0.35, l2, size=12, color=P["on_dark"])
            text(s, 0.6, 6.95, 8.8, 0.3, "Department of Computer Engineering  ·  Bajaj Institute of Technology, Wardha",
                 size=10, color=P["on_dark"])
        return s

    def agenda(self, items):
        P = self.P
        s = self.content("Agenda", "What we will cover")
        half = (len(items) + 1) // 2
        for i, it in enumerate(items):
            col, row = divmod(i, half)
            x = 0.6 + col * 4.55
            y = 1.75 + row * 0.82
            text(s, x, y, 0.75, 0.6, f"{i + 1:02d}", size=22, bold=True, font=HEAD, color=P["primary"], anchor="m")
            text(s, x + 0.8, y, 3.5, 0.6, it, size=16, color=P["text"], anchor="m")
            if row < half - 1 and not (col == 1 and i == len(items) - 1):
                line(s, x, y + 0.71, x + 4.2, y + 0.71, P["line"], 0.75)
        return s

    # -------------------------------------------- reusable compositions
    def subtitle(self, s, t):
        text(s, 0.6, 1.38, 8.8, 0.32, t, size=13, italic=True, color=self.P["muted"])

    def numbered_rows(self, s, rows, y0=1.72, step=1.02):
        P = self.P
        for i, (head, desc, ic) in enumerate(rows):
            y = y0 + i * step
            text(s, 0.6, y, 0.9, 0.82, f"{i + 1:02d}", size=30, bold=True, font=HEAD, color=P["primary"], anchor="m")
            text(s, 1.6, y + 0.06, 7.1, 0.36, head, size=16, bold=True, color=P["ink"])
            text(s, 1.6, y + 0.42, 7.1, 0.45, desc, size=14, color=P["muted"])
            self.icon(s, ic, 8.95, y + 0.2, 0.42, P["ic_p"])
            if i < len(rows) - 1:
                line(s, 0.6, y + step - 0.08, 9.4, y + step - 0.08, P["line"], 0.75)

    def insight(self, s, y, label, body, h=0.72, icon="lightbulb"):
        P = self.P
        box(s, 0.6, y, 8.8, h, bg=P["ink"], radius=0.1)
        self.icon(s, icon, 0.82, y + (h - 0.36) / 2, 0.36, P["ic_a"])
        text(s, 1.35, y, 7.9, h, [[(label + "  ", {"bold": True, "color": P["accent"]}), (body, {})]],
             size=13.5, color=WHITE, anchor="m")

    def theme_card(self, s, x, y, w, h, ic, title, refs, body, alt=False, size=13):
        P = self.P
        box(s, x, y, w, h, bg=P["card"], radius=0.12)
        self.icon_circle(s, ic, x + 0.2, y + 0.2, 0.5, bg=P["accent"] if alt else P["primary"])
        text(s, x + 0.85, y + 0.17, w - 1.0, 0.56,
             [[(title, {"bold": True, "color": P["ink"], "size": 14.5})],
              [(refs, {"color": P["accent"] if alt else P["primary"], "size": 11, "bold": True})]],
             anchor="m", after=0)
        text(s, x + 0.2, y + 0.85, w - 0.4, h - 0.92, body, size=size, color=P["text"])

    def lit_table(self, s, rows, row_h, widths=(0.6, 3.7, 4.5)):
        P = self.P
        nr = len(rows) + 1
        gf = s.shapes.add_table(nr, 3, Inches(0.6), Inches(1.6), Inches(sum(widths)), Inches(0.45 + row_h * len(rows)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for k in ("bandRow", "firstRow"):
            tblPr.set(k, "0")
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        head = ["#", "Paper", "Work reported"]
        for i in range(nr):
            tbl.rows[i].height = Inches(0.45 if i == 0 else row_h)
            for j in range(3):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.1)
                c.margin_top = c.margin_bottom = Inches(0.05)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(P["ink"] if i == 0 else (WHITE if i % 2 else P["card"]))
                tf = c.text_frame
                tf.word_wrap = True
                if i == 0:
                    fill(tf, head[j], 11, WHITE, True, spc=100)
                    continue
                n, t, v, w = rows[i - 1]
                if j == 0:
                    fill(tf, n, 15, P["primary"], True, font=HEAD)
                elif j == 1:
                    fill(tf, [[(t, {"bold": True})], [(v, {"size": 11, "color": P["muted"]})]], 12.5, P["ink"], after=2)
                else:
                    fill(tf, w, 12.5, P["text"])
        return tbl

    def baseline(self, s, primary, supporting):
        P = self.P
        tag, ref, t, v, why, used = primary
        box(s, 0.6, 1.65, 4.3, 5.1, bg=P["ink"], radius=0.14)
        text(s, 0.9, 1.95, 3.8, 0.3, tag, size=10.5, bold=True, color=P["accent"], spc=180)
        text(s, 0.9, 2.3, 3.8, 0.7, ref, size=34, bold=True, font=HEAD, color=P["accent"])
        text(s, 0.9, 3.0, 3.75, 1.2, t, size=17, bold=True, font=HEAD, color=WHITE)
        text(s, 0.9, 4.15, 3.75, 0.3, v, size=12, italic=True, color=P["on_dark"])
        for k, (lab, body) in enumerate((("WHY SELECTED", why), ("USED FOR", used))):
            yy = 4.65 + k * 1.0
            text(s, 0.9, yy, 3.75, 0.25, lab, size=9.5, bold=True, color=P["accent"], spc=160)
            text(s, 0.9, yy + 0.27, 3.75, 0.7, body, size=13, color=WHITE)
        for k, (tag, ref, t, v, why, used) in enumerate(supporting):
            y = 1.65 + k * 2.6
            box(s, 5.1, y, 4.3, 2.45, bg=P["card"], radius=0.14)
            text(s, 5.35, y + 0.18, 2.5, 0.25, tag, size=9.5, bold=True, color=P["primary"], spc=160)
            text(s, 8.4, y + 0.08, 0.8, 0.5, ref, size=20, bold=True, font=HEAD, color=P["primary"], align="r")
            text(s, 5.35, y + 0.48, 3.8, 0.62, t, size=13.5, bold=True, color=P["ink"])
            text(s, 5.35, y + 1.1, 3.8, 0.26, v, size=11, italic=True, color=P["muted"])
            text(s, 5.35, y + 1.42, 3.8, 0.95, [[("Why: ", {"bold": True, "color": P["primary"]}), (why, {})],
                                                [("Used for: ", {"bold": True, "color": P["primary"]}), (used, {})]],
                 size=12, color=P["text"], after=3)

    def code(self, s, x, y, w, h, lines, size=11.5, title="PSEUDOCODE"):
        P = self.P
        box(s, x, y, w, h, bg=P["ink"], radius=0.12)
        text(s, x + 0.22, y + 0.16, w - 0.4, 0.26, title, size=9.5, bold=True, color=P["accent"], spc=180)
        paras = []
        for i, ln in enumerate(lines, 1):
            code_part, _, comment = ln.partition("#")
            runs = [(f"{i:>2}  ", {"color": P["code_num"]}),
                    (code_part, {"color": "F1F0FA"})]
            if comment:
                runs.append(("# " + comment.strip(), {"color": P["on_dark"], "italic": True}))
            paras.append(runs)
        text(s, x + 0.2, y + 0.5, w - 0.35, h - 0.6, paras, size=size, font=MONO, after=2)

    def three_columns(self, s, cols, y=1.65, h=5.15):
        P = self.P
        for k, (head, col, ic, items) in enumerate(cols):
            x = 0.6 + k * 3.025
            w = 2.75
            box(s, x, y, w, h, bg=P["card"], radius=0.12)
            box(s, x + 0.2, y + 0.2, 0.46, 0.46, bg=col, shape="oval")
            self.icon(s, ic, x + 0.31, y + 0.31, 0.24, "w")
            text(s, x + 0.78, y + 0.18, w - 0.9, 0.5, head, size=16, bold=True, color=col, anchor="m", font=HEAD)
            for i, (lead, rest) in enumerate(items):
                yy = y + 0.9 + i * 0.84
                text(s, x + 0.2, yy, w - 0.35, 0.78, [[(lead, {"bold": True, "color": P["ink"]})],
                                                      [(rest, {"size": 12, "color": P["muted"]})]],
                     size=14, after=1, anchor="m")
                if i < len(items) - 1:
                    line(s, x + 0.2, yy + 0.81, x + w - 0.2, yy + 0.81, P["line"], 0.75)

    def app_grid(self, s, apps):
        P = self.P
        for i, (ic, t, d) in enumerate(apps):
            x = 0.6 + (i % 2) * 4.5
            y = 1.65 + (i // 2) * 1.72
            box(s, x, y, 4.3, 1.55, bg=P["card"], radius=0.12)
            self.icon_circle(s, ic, x + 0.25, y + 0.4, 0.75, bg=P["accent"] if i in (1, 4) else P["primary"])
            text(s, x + 1.25, y + 0.2, 2.9, 0.42, t, size=16, bold=True, color=P["ink"], font=HEAD, anchor="m")
            text(s, x + 1.25, y + 0.64, 2.9, 0.85, d, size=13, color=P["muted"])

    def comparison_table(self, s, rows, widths, y=1.65, row_h=0.52, hl_col=None):
        P = self.P
        nr, nc = len(rows), len(widths)
        gf = s.shapes.add_table(nr, nc, Inches(0.6), Inches(y), Inches(sum(widths)), Inches(row_h * nr))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for k in ("bandRow", "firstRow"):
            tblPr.set(k, "0")
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(nr):
            tbl.rows[i].height = Inches(row_h)
            for j in range(nc):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.1)
                c.margin_top = c.margin_bottom = Inches(0.04)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid()
                if i == 0:
                    bgc = P["ink"] if j != hl_col else P["primary"]
                elif j == hl_col:
                    bgc = P["soft"]
                else:
                    bgc = WHITE if i % 2 else P["card"]
                c.fill.fore_color.rgb = rgb(bgc)
                tf = c.text_frame
                tf.word_wrap = True
                if i == 0:
                    fill(tf, rows[i][j], 12, WHITE, True)
                else:
                    fill(tf, rows[i][j], 12.5, P["ink"] if j == 0 else P["text"], j == 0)
        return tbl

    def conclusion(self, title, takeaways, future, statement):
        P = self.P
        s = self.content("Conclusion & future scope", title, dark=True)
        for k, (lab, items, ic) in enumerate((("KEY TAKEAWAYS", takeaways, "check"),
                                              ("FUTURE SCOPE", future, "arrowr"))):
            x = 0.6 + k * 4.55
            text(s, x, 1.65, 4.2, 0.3, lab, size=10.5, bold=True, color=P["accent"], spc=180)
            for i, it in enumerate(items):
                y = 2.05 + i * 0.9
                self.icon(s, ic, x, y + 0.2, 0.3, P["ic_a"] if k else "w")
                text(s, x + 0.48, y, 3.75, 0.78, it, size=14.5, color=WHITE, anchor="m")
                if i < len(items) - 1:
                    line(s, x, y + 0.85, x + 4.2, y + 0.85, P["dark_line"], 0.75)
        box(s, 0.6, 5.8, 8.8, 0.95, bg=P["primary"], radius=0.12, paras=statement, size=16, bold=True,
            italic=True, color=WHITE, font=HEAD, margin=0.2)
        return s

    def references(self, title, refs, start=1, size=14):
        P = self.P
        s = self.content("References", title)
        paras = []
        for i, r in enumerate(refs, start):
            paras.append({"runs": [(f"[{i}]\t", {"bold": True, "color": P["primary"]}), (r, {})]})
        tb = text(s, 0.6, 1.65, 8.8, 5.2, paras, size=size, color=P["text"], after=12)
        for p in tb.text_frame.paragraphs:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.45))))
            pPr.set("indent", str(-int(Inches(0.45))))
        return s

    def closing(self, motif, name):
        P = self.P
        s = self.cover("Seminar-I  ·  2025–26", "Thank you", None, motif, title_size=54)
        text(s, 0.6, 3.45, 6.0, 0.6, "Questions & discussion", size=22, italic=True, font=HEAD, color=P["on_dark"])
        text(s, 0.6, 5.75, 6.0, 0.45, name, size=17, bold=True, font=HEAD, color=WHITE)
        text(s, 0.6, 6.2, 8.0, 0.35, "Department of Computer Engineering  ·  Bajaj Institute of Technology, Wardha",
             size=11, color=P["on_dark"])
        return s
