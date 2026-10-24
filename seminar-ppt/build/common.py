"""Shared helpers for building the seminar decks on top of the BIT Seminar-I template.

Every slide is a clone of one of the template's own slides, so the template's
layouts, title styling, logo, header banner and slide numbers are preserved.
Diagrams are drawn as native (editable) PowerPoint shapes.
"""
import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Template (Office theme) palette
NAVY = "1F497D"
BLUE = "4F81BD"
LIGHT = "DCE6F2"
PALE = "F2F6FB"
TEAL = "31859C"
ORANGE = "E46C0A"
RED = "C0504D"
GREEN = "76923C"
GOLD = "F2C314"
INK = "262626"
MUTED = "595959"
WHITE = "FFFFFF"
CODE_BG = "F2F2F2"

TITLE_FONT = "Times New Roman"
BODY_FONT = "Calibri"
MONO_FONT = "Courier New"

# Template slide numbers (1-based) used as frames
T_TITLE, T_INDEX, T_CONTENT, T_CONCLUSION, T_REFERENCES = 3, 4, 6, 14, 15

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, template, icon_dir):
        self.prs = Presentation(template)
        self.src = list(self.prs.slides)
        self.n_orig = len(self.src)
        self.icon_dir = icon_dir
        self.logo = self._extract_logo()

    def _extract_logo(self):
        # The round BIT logo used on the template's Conclusion slide
        for shp in self.src[T_CONCLUSION - 1].shapes:
            if shp.shape_type == 13:
                return shp.image
        raise RuntimeError("logo not found")

    # ---------- slide cloning ----------
    def clone(self, idx):
        src = self.src[idx - 1]
        new = self.prs.slides.add_slide(src.slide_layout)
        tree = new.shapes._spTree
        keep = (qn("p:nvGrpSpPr"), qn("p:grpSpPr"))
        for el in list(tree):
            if el.tag not in keep:
                tree.remove(el)
        rmap = {}
        for rId, rel in src.part.rels.items():
            if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
                continue
            if rel.is_external:
                rmap[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rmap[rId] = new.part.relate_to(rel.target_part, rel.reltype)
        for el in src.shapes._spTree:
            if el.tag in keep:
                continue
            e = copy.deepcopy(el)
            for node in e.iter():
                for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                    v = node.get(attr)
                    if v in rmap:
                        node.set(attr, rmap[v])
            tree.append(e)
        return new

    def save(self, path):
        lst = self.prs.slides._sldIdLst
        for sldId in list(lst)[: self.n_orig]:
            self.prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
        self.prs.save(path)

    # ---------- frames ----------
    def content_slide(self, title, size=30, frame=T_CONTENT):
        """Template content frame: centred Times New Roman title, logo top-left, slide number."""
        s = self.clone(frame)
        for shp in list(s.shapes):
            if shp.is_placeholder and shp.placeholder_format.idx == 1:
                shp._element.getparent().remove(shp._element)
            elif shp.shape_type == 13:  # template's own logo; re-added at the standard spot
                shp._element.getparent().remove(shp._element)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(1.3), Inches(0.2), Inches(7.4), Inches(0.9)
        tf = t.text_frame
        tf.word_wrap = True
        bp = tf._txBody.find(qn("a:bodyPr"))
        for ch in list(bp):
            bp.remove(ch)
        bp.set("anchor", "ctr")
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = tf.paragraphs[0]
        for r in list(p._p):
            if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                p._p.remove(r)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.size = Pt(size)
        run.font.name = TITLE_FONT
        run.font.color.rgb = rgb("000000")
        self.add_logo(s)
        return s

    def add_logo(self, s, x=0.18, y=0.14, w=1.08):
        import io

        h = w * self.logo.size[1] / self.logo.size[0]
        s.shapes.add_picture(io.BytesIO(self.logo.blob), Inches(x), Inches(y), Inches(w), Inches(h))

    def icon(self, s, name, x, y, d, color="w"):
        path = os.path.join(self.icon_dir, f"{name}_{color}.png")
        return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(d), Inches(d))

    def icon_circle(self, s, name, x, y, d, fill=NAVY, color="w"):
        box(s, x, y, d, d, fill=fill, shape="oval")
        pad = d * 0.24
        return self.icon(s, name, x + pad, y + pad, d - 2 * pad, color)


# ---------- text ----------
def _apply_run(run, size, color, bold, italic, font):
    f = run.font
    f.size = Pt(size)
    f.color.rgb = rgb(color)
    f.bold = bold
    f.italic = italic
    f.name = font


def _bullet(p, char="•", color=None, indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if color:
        bc = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(bc, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def fill_text(tf, paras, size=14, color=INK, bold=False, italic=False, align="l",
              font=BODY_FONT, after=4, bullet=None, line=None):
    """paras: list of items; each item is a str, or a list of runs (text, {opts}),
    or a dict {"runs": ..., "bullet": bool, "size":.., "align":.., "after":..}."""
    first = True
    for item in paras:
        opts = {}
        if isinstance(item, dict):
            opts = item
            runs = item["runs"]
        else:
            runs = item
        if isinstance(runs, str):
            runs = [(runs, {})]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[opts.get("align", align)]
        p.space_after = Pt(opts.get("after", after))
        if line:
            p.line_spacing = line
        b = opts.get("bullet", bullet)
        if b:
            _bullet(p, char=b if isinstance(b, str) else "•", color=opts.get("bullet_color"))
        for text, ro in runs:
            r = p.add_run()
            r.text = text
            _apply_run(r, ro.get("size", opts.get("size", size)), ro.get("color", opts.get("color", color)),
                       ro.get("bold", opts.get("bold", bold)), ro.get("italic", italic),
                       ro.get("font", font))


def text(s, x, y, w, h, paras, size=14, color=INK, bold=False, italic=False, align="l",
         anchor="t", font=BODY_FONT, margin=0.04, after=4, bullet=None, line=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = ANCHOR[anchor]
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if isinstance(paras, str):
        paras = [paras]
    fill_text(tf, paras, size, color, bold, italic, align, font, after, bullet, line)
    return tb


# ---------- shapes ----------
SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rrect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.CHEVRON,
    "rarrow": MSO_SHAPE.RIGHT_ARROW,
    "darrow": MSO_SHAPE.DOWN_ARROW,
    "callout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
}


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(s, x, y, w, h, fill=None, line=None, lw=1.0, shape="rrect", radius=0.08,
        paras=None, size=13, color=INK, bold=False, italic=False, align="c", anchor="m",
        font=BODY_FONT, margin=0.06, after=0, dash=False):
    shp = s.shapes.add_shape(SHAPES[shape], Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(shp)
    if shape == "rrect":
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = 7  # dash
    else:
        shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = ANCHOR[anchor]
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if paras is not None:
        if isinstance(paras, str):
            paras = [paras]
        fill_text(tf, paras, size, color, bold, italic, align, font, after)
    return shp


def _line(s, x1, y1, x2, y2, color, w):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    return c


def _head(conn, which):
    ln = conn.line._get_or_add_ln()
    etree.SubElement(ln, qn(f"a:{which}"), type="triangle", w="med", len="med")


def arrow(s, x1, y1, x2, y2, color=MUTED, w=1.75, head="end", dash=False):
    c = _line(s, x1, y1, x2, y2, color, w)
    if dash:
        c.line.dash_style = 7
    if head in ("start", "both"):
        _head(c, "headEnd")
    if head in ("end", "both"):
        _head(c, "tailEnd")
    return c


def polyline(s, pts, color=MUTED, w=1.75, head="end", dash=False):
    """Orthogonal path through points; arrowhead on the final segment."""
    n = len(pts) - 1
    for i in range(n):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        h = None
        if i == n - 1 and head in ("end", "both"):
            h = "end"
        if i == 0 and head in ("start", "both"):
            h = "both" if h else "start"
        arrow(s, x1, y1, x2, y2, color, w, head=h, dash=dash)


def line(s, x1, y1, x2, y2, color=MUTED, w=1.25):
    return _line(s, x1, y1, x2, y2, color, w)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def table(s, x, y, col_w, rows, size=12, header_size=13, row_h=0.4, first_col_bold=False,
          col_align=None, fills=None, header_h=0.5):
    nr, nc = len(rows), len(col_w)
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gf.table
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nr):
        tbl.rows[i].height = Inches(header_h if i == 0 else row_h)
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[i][j]
            tf = cell.text_frame
            tf.word_wrap = True
            if i == 0:
                fill_text(tf, [val] if isinstance(val, str) else val, header_size, WHITE, True,
                          align="l", after=0)
            else:
                if isinstance(val, str):
                    val = [val]
                fill_text(tf, val, size, INK, first_col_bold and j == 0,
                          align=(col_align[j] if col_align else "l"), after=0)
            if fills and i > 0 and fills.get((i, j)):
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fills[(i, j)])
    return tbl


# ---------- reusable compositions ----------
def title_slide(deck, title_lines, student, roll, title_size=24):
    s = deck.clone(T_TITLE)
    ph = [p for p in s.placeholders][0]
    tf = ph.text_frame
    body = tf._txBody
    for p in body.findall(qn("a:p")):
        body.remove(p)
    bp = body.find(qn("a:bodyPr"))
    for ch in list(bp):
        bp.remove(ch)
    etree.SubElement(bp, qn("a:noAutofit"))
    bp.set("anchor", "t")
    ph.top = Inches(1.55)
    ph.height = Inches(5.75)

    def para(runs, size, after=0, bold=True):
        p = etree.SubElement(body, qn("a:p"))
        pPr = etree.SubElement(p, qn("a:pPr"), algn="ctr")
        sa = etree.SubElement(pPr, qn("a:spcAft"))
        etree.SubElement(sa, qn("a:spcPts"), val=str(int(after * 100)))
        for t, color in runs:
            r = etree.SubElement(p, qn("a:r"))
            rPr = etree.SubElement(r, qn("a:rPr"), lang="en-US", sz=str(int(size * 100)),
                                   b="1" if bold else "0", dirty="0")
            sf = etree.SubElement(rPr, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
            etree.SubElement(rPr, qn("a:latin"), typeface=TITLE_FONT)
            etree.SubElement(rPr, qn("a:cs"), typeface=TITLE_FONT)
            etree.SubElement(r, qn("a:t")).text = t

    for i, tl in enumerate(title_lines):
        para([(tl, NAVY)], title_size, after=14 if i == len(title_lines) - 1 else 0)
    para([("Delivered By:", "000000")], 20, after=4)
    para([(student, "000000")], 18)
    para([(roll, "000000")], 16, after=12, bold=False)
    para([("Department of Computer Engineering", "000000")], 18)
    para([("Bajaj Institute of Technology, Wardha.", "000000")], 18, after=12)
    para([("Under Supervision of:", "000000")], 18, after=4)
    para([("Name of Faculty", "FF0000")], 18)
    para([("Designation of faculty", "FF0000")], 18)
    para([("Computer Engineering Department, BIT, Wardha", "000000")], 18)
    return s


def thank_you_slide(deck, student):
    s = deck.clone(T_TITLE)
    ph = [p for p in s.placeholders][0]
    ph._element.getparent().remove(ph._element)
    text(s, 0.75, 2.5, 8.5, 1.1, "Thank You", size=48, bold=True, color=NAVY, align="c",
         font=TITLE_FONT, anchor="m")
    text(s, 0.75, 3.6, 8.5, 0.6, "Questions & Discussion", size=24, color=INK, align="c",
         font=TITLE_FONT, anchor="m")
    text(s, 0.75, 5.0, 8.5, 1.1, [student, "Department of Computer Engineering, BIT, Wardha"],
         size=16, color=MUTED, align="c", font=TITLE_FONT, after=2)
    return s


def index_slide(deck, items):
    s = deck.clone(T_INDEX)
    for shp in s.shapes:
        if shp.name == "TextBox 7":
            tb = shp
    txBody = tb.text_frame._txBody
    ps = txBody.findall(qn("a:p"))
    proto = copy.deepcopy(ps[1])
    for p in ps:
        txBody.remove(p)
    for it in items:
        p = copy.deepcopy(proto)
        r = p.find(qn("a:r"))
        r.find(qn("a:rPr")).set("sz", "2200")
        r.find(qn("a:t")).text = "  " + it
        sa = etree.SubElement(p.find(qn("a:pPr")), qn("a:spcBef"))
        etree.SubElement(sa, qn("a:spcPts"), val="300")
        txBody.append(p)
    return s


def references_slide(deck, refs, start=1, size=14):
    s = deck.clone(T_REFERENCES)
    body = None
    for shp in s.placeholders:
        if shp.placeholder_format.idx != 0 and shp.placeholder_format.type != 13:
            body = shp
    tf = body.text_frame
    txBody = tf._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    bp = txBody.find(qn("a:bodyPr"))
    for ch in list(bp):
        bp.remove(ch)
    etree.SubElement(bp, qn("a:noAutofit"))
    body.left, body.top, body.width, body.height = Inches(0.35), Inches(1.4), Inches(9.3), Inches(5.45)
    etree.SubElement(txBody, qn("a:p"))
    paras = []
    for i, ref in enumerate(refs, start):
        runs = [(f"[{i}]  ", {"bold": True, "color": NAVY})]
        runs += ref if isinstance(ref, list) else [(ref, {})]
        paras.append({"runs": runs})
    fill_text(tf, paras, size=size, color=INK, font=TITLE_FONT, after=10)
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.45))))
        pPr.set("indent", str(-int(Inches(0.45))))
        etree.SubElement(pPr, qn("a:buNone"))
    return s


def numbered_rows(deck, s, rows, y0=1.35, step=1.08, h=0.92, fill=PALE, circ=NAVY, icon_color="b"):
    """Objectives-style rows: number badge, bold lead + description, icon on the right."""
    for i, (lead, desc, ic) in enumerate(rows):
        y = y0 + i * step
        box(s, 0.5, y, 9.0, h, fill=fill)
        box(s, 0.68, y + (h - 0.56) / 2, 0.56, 0.56, fill=circ, shape="oval", paras=str(i + 1),
            size=20, bold=True, color=WHITE)
        text(s, 1.42, y + 0.04, 7.2, h - 0.08,
             [[(lead + ": ", {"bold": True, "color": NAVY}), (desc, {})]], size=16, anchor="m")
        deck.icon(s, ic, 8.82, y + (h - 0.5) / 2, 0.5, icon_color)


def theme_card(deck, s, x, y, w, h, icon, title, refs, body, fill=PALE, circ=NAVY, size=14):
    box(s, x, y, w, h, fill=fill)
    deck.icon_circle(s, icon, x + 0.15, y + 0.15, 0.55, fill=circ)
    text(s, x + 0.8, y + 0.12, w - 0.9, 0.62,
         [[(title, {"bold": True, "color": NAVY, "size": 14})],
          [(refs, {"color": MUTED, "size": 11})]], anchor="m", after=0)
    text(s, x + 0.15, y + 0.82, w - 0.3, h - 0.9, body, size=size, color=INK)


def three_columns(deck, s, cols, y=1.35, h=5.4, w=2.85, xs=(0.5, 3.575, 6.65), item_h=0.86):
    """Features / Advantages / Disadvantages columns: (header, color, icon, items)."""
    for x, (head, col, ic, items) in zip(xs, cols):
        box(s, x, y, w, h, fill=PALE)
        box(s, x, y, w, 0.62, fill=col, paras=[[(head, {"bold": True})]], size=17, color=WHITE,
            align="l", margin=0.12)
        deck.icon(s, ic, x + w - 0.52, y + 0.11, 0.4, "w")
        for i, (lead, rest) in enumerate(items):
            yy = y + 0.8 + i * (item_h + 0.08)
            box(s, x + 0.12, yy, w - 0.24, item_h, fill=WHITE, line=LIGHT, lw=0.75)
            text(s, x + 0.22, yy + 0.03, w - 0.44, item_h - 0.06,
                 [[(lead, {"bold": True, "color": col})], [(rest, {"size": 12, "color": INK})]],
                 size=14, anchor="m", after=0)


def app_grid(deck, s, apps, y0=1.45, h=2.45, gap_y=0.3, circ=NAVY):
    xs = (0.5, 3.575, 6.65)
    w = 2.85
    for i, (ic, title, desc) in enumerate(apps):
        x = xs[i % 3]
        y = y0 + (i // 3) * (h + gap_y)
        box(s, x, y, w, h, fill=PALE)
        deck.icon_circle(s, ic, x + (w - 0.8) / 2, y + 0.2, 0.8, fill=circ)
        text(s, x + 0.12, y + 1.1, w - 0.24, 0.42, title, size=16, bold=True, color=NAVY,
             align="c", anchor="m")
        text(s, x + 0.15, y + 1.52, w - 0.3, h - 1.6, desc, size=14, align="c")


def code_box(s, x, y, w, h, lines, title=None, size=11.5):
    box(s, x, y, w, h, fill=CODE_BG, line="BFBFBF", lw=0.75, radius=0.06)
    yy = y + 0.08
    if title:
        text(s, x + 0.12, yy, w - 0.24, 0.32, title, size=13, bold=True, color=NAVY)
        yy += 0.36
    paras = []
    for i, ln in enumerate(lines, 1):
        paras.append([(f"{i:>2}  ", {"color": "7F7F7F"}), (ln, {})])
    text(s, x + 0.1, yy, w - 0.2, h - (yy - y) - 0.05, paras, size=size, font=MONO_FONT, after=1)


def conclusion_slide(deck, takeaways, future, banner):
    s = deck.content_slide("Conclusion", frame=T_CONCLUSION)
    for x, head, items, ic, col in ((0.5, "Key Takeaways", takeaways, "check", NAVY),
                                    (5.1, "Future Scope", future, "lightbulb", ORANGE)):
        deck.icon_circle(s, ic, x, 1.35, 0.5, fill=col)
        text(s, x + 0.6, 1.35, 3.6, 0.5, head, size=20, bold=True, color=col, anchor="m")
        for i, it in enumerate(items):
            yy = 2.05 + i * 0.95
            box(s, x, yy, 4.4, 0.82, fill=PALE)
            text(s, x + 0.15, yy + 0.02, 4.1, 0.78, it, size=14, anchor="m")
    box(s, 0.5, 5.95, 9.0, 0.82, fill=NAVY, paras=banner, size=15, bold=True, color=WHITE,
        margin=0.12)
    return s
