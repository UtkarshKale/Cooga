"""Vedika's design system: "hazard signage".

Deliberately independent of design.py's look: all-sans Arial, white slides, sharp-cornered
blocks, amber + charcoal signage colours, numbered section tiles beside plain section titles,
logo and page tile in the footer, light pseudocode panels. Only the shape/text primitives
are shared with design.py.
"""
import io
import os

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

from design import MONO, arrow, box, fill, line, notes, path, rgb, text  # noqa: F401

FONT = "Arial"
CHAR = "23272B"
AMBER = "F4A81D"
AMBER_LT = "FDF1D8"
RED = "D9472B"
BLUE = "2C6E91"
GREEN = "3C9D5D"
GRAY = "F0F2F4"
GRAY2 = "D5DADF"
MUTED = "5F6B76"
WHITE = "FFFFFF"
ICON = {CHAR: "vc", AMBER: "va", RED: "vr", BLUE: "vb", GREEN: "vg", WHITE: "w"}


def t(s, x, y, w, h, paras, size=13, color=CHAR, **kw):
    kw.setdefault("font", FONT)
    return text(s, x, y, w, h, paras, size=size, color=color, **kw)


def b(s, x, y, w, h, bg=None, paras=None, size=12, color=CHAR, **kw):
    kw.setdefault("font", FONT)
    kw.setdefault("shape", "rect")
    return box(s, x, y, w, h, bg=bg, paras=paras, size=size, color=color, **kw)


class VDeck:
    def __init__(self, template, icon_dir, footer):
        self.icon_dir = icon_dir
        self.footer = footer
        self.prs = Presentation(template)
        self.logo = self.banner = None
        for sl in self.prs.slides:
            for shp in sl.shapes:
                if shp.shape_type == 13:
                    if shp.image.size == (1100, 140):
                        self.banner = shp.image.blob
                    elif shp.image.size == (316, 238):
                        self.logo = shp.image
        lst = self.prs.slides._sldIdLst
        for sldId in list(lst):
            self.prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
        self.blank = self.prs.slide_masters[0].slide_layouts[6]

    def save(self, p):
        self.prs.save(p)

    # ------------------------------------------------ primitives
    def icon(self, s, name, x, y, d, color):
        return s.shapes.add_picture(os.path.join(self.icon_dir, f"{name}_{ICON[color]}.png"),
                                    Inches(x), Inches(y), Inches(d), Inches(d))

    def icon_tile(self, s, name, x, y, d, bg=AMBER, fg=CHAR):
        b(s, x, y, d, d, bg=bg)
        pad = d * 0.22
        self.icon(s, name, x + pad, y + pad, d - 2 * pad, fg)

    def logo_pic(self, s, x, y, w):
        h = w * self.logo.size[1] / self.logo.size[0]
        s.shapes.add_picture(io.BytesIO(self.logo.blob), Inches(x), Inches(y), Inches(w), Inches(h))

    def _slide(self):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(WHITE)
        return s

    def _page_tile(self, s):
        shp = b(s, 9.05, 6.95, 0.36, 0.36, bg=AMBER, margin=0)
        p = shp.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        fld = etree.SubElement(p._p, qn("a:fld"), id="{3C1D2E4F-6A7B-4C8D-9E0F-1A2B3C4D5E6F}", type="slidenum")
        rPr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz="1000", b="1")
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=CHAR)
        etree.SubElement(rPr, qn("a:latin"), typeface=FONT)
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))

    # ------------------------------------------------ frames
    def slide(self, num, title, sub=None):
        s = self._slide()
        b(s, 0.6, 0.42, 0.62, 0.62, bg=AMBER, paras=num, size=18, bold=True, margin=0)
        t(s, 1.4, 0.36, 8.0, 0.48, title, size=24, bold=True, anchor="m")
        if sub:
            t(s, 1.4, 0.82, 8.0, 0.32, sub, size=13, color=MUTED)
        self.logo_pic(s, 0.6, 6.93, 0.48)
        t(s, 1.18, 6.98, 6.5, 0.3, self.footer, size=9, color=MUTED, anchor="m")
        self._page_tile(s)
        return s

    def cover(self, title, kicker, hazard_icons, info=None, title_size=40, extra=None, extra_y=4.3):
        s = self._slide()
        self.logo_pic(s, 0.35, 0.22, 1.25)
        pic = s.shapes.add_picture(io.BytesIO(self.banner), Inches(1.8), Inches(0.2), Inches(7.85), Inches(1.17))
        pic.crop_left, pic.crop_right = 0.11745, 0.03185
        b(s, 0, 1.75, 6.6, 3.25, bg=AMBER)
        t(s, 0.6, 2.0, 5.6, 0.32, kicker, size=12, bold=True, spc=120)
        t(s, 0.6, 2.4, 5.7, 1.9, title, size=title_size, bold=True, line=0.95)
        if extra:
            t(s, 0.6, extra_y, 5.7, 0.5, extra, size=18)
        b(s, 6.6, 1.75, 3.4, 3.25, bg=CHAR)
        for i, (ic, lab) in enumerate(hazard_icons):
            cx = 7.55 + (i % 2) * 1.5
            cy = 2.15 + (i // 2) * 1.45
            self.icon(s, ic, cx - 0.33, cy, 0.66, AMBER)
            t(s, cx - 0.7, cy + 0.72, 1.4, 0.3, lab, size=11, color=WHITE, align="c", bold=True)
        if info:
            for k, (lab, l1, l2) in enumerate(info):
                x = 0.6 + k * 4.5
                t(s, x, 5.4, 4.2, 0.28, lab.upper(), size=10.5, bold=True, color=MUTED, spc=120)
                t(s, x, 5.7, 4.2, 0.45, l1, size=20, bold=True)
                t(s, x, 6.15, 4.2, 0.35, l2, size=12.5, color=MUTED)
            t(s, 0.6, 6.85, 8.8, 0.3, "Department of Computer Engineering, Bajaj Institute of Technology, Wardha",
              size=10.5, color=MUTED)
        return s

    def code(self, s, x, y, w, h, lines, size=11.5):
        b(s, x, y, w, h, bg=GRAY)
        b(s, x, y, 1.3, 0.3, bg=CHAR, paras="PSEUDOCODE", size=9, bold=True, color=AMBER, spc=100, margin=0)
        paras = []
        for i, ln in enumerate(lines, 1):
            code_part, _, comment = ln.partition("#")
            runs = [(f"{i:>2}  ", {"color": "8B96A0"}), (code_part, {"color": CHAR})]
            if comment:
                runs.append(("# " + comment.strip(), {"color": BLUE, "italic": True}))
            paras.append(runs)
        text(s, x + 0.2, y + 0.45, w - 0.3, h - 0.5, paras, size=size, font=MONO, after=2)

    def key_block(self, s, y, label, body, h=0.75, dark=False):
        b(s, 0.6, y, 8.8, h, bg=CHAR if dark else AMBER)
        t(s, 0.85, y, 8.4, h, [[(label + ":  ", {"bold": True, "color": AMBER if dark else CHAR}), (body, {})]],
          size=13.5, color=WHITE if dark else CHAR, anchor="m")

    def table(self, s, rows, widths, y=1.5, header_h=0.45, row_h=0.6, hl_col=None, size=12, num_col=False):
        nr, nc = len(rows), len(widths)
        gf = s.shapes.add_table(nr, nc, Inches(0.6), Inches(y), Inches(sum(widths)),
                                Inches(header_h + row_h * (nr - 1)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for k in ("bandRow", "firstRow"):
            tblPr.set(k, "0")
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(nr):
            tbl.rows[i].height = Inches(header_h if i == 0 else row_h)
            for j in range(nc):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.1)
                c.margin_top = c.margin_bottom = Inches(0.05)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid()
                if i == 0:
                    bg = AMBER
                elif j == hl_col:
                    bg = AMBER_LT
                else:
                    bg = WHITE if i % 2 else GRAY
                c.fill.fore_color.rgb = rgb(bg)
                tf = c.text_frame
                tf.word_wrap = True
                v = rows[i][j]
                if i == 0:
                    fill(tf, v, 12, CHAR, True, font=FONT)
                elif num_col and j == 0:
                    fill(tf, v, 14, CHAR, True, font=FONT, align="c")
                else:
                    fill(tf, v, size, CHAR, j == 0 and not num_col, font=FONT, after=2)
        return tbl
